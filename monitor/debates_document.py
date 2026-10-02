"""The Word document attached to the debate summaries email.

WHY IT EXISTS
-------------
2 October 2026: the directorate asked for the debate emails to work like the
monitoring company's committee notes — "a short summary in the email body and
another word document that outlines the session in more detail". Their note
of the housing committee's 1 October meeting gave, for each session, the
witnesses and their posts, the evidence papers, and every contribution in two
to four sentences, in order.

So the email keeps its one sentence per speaker, and this document gives,
for every item in the email:

  * the meeting, with links to the Record, Senedd.tv and the agenda papers;
  * for a committee, the witnesses who gave evidence, with their posts (from
    the Record's "Others in Attendance");
  * every contribution, in order, with the fuller account Claude writes for
    it (``Point.detail``) — or, without the key, the speaker's own sentences.

The evidence papers themselves are not listed one by one: the Senedd's
business site, where they are published, refuses requests from the server
this runs on. The document links to the meeting's agenda page, which lists
them.
"""

from __future__ import annotations

import io
import re
from datetime import date

from .weekly_document import (DARK_BLUE, MUTED, OFF_BLACK, ORANGE, _Doc,
                              _border_bottom, _border_left, _fix_settings,
                              _hyperlink, _page_number, _pt, _rgb, tidy)


def _day(d: date | None) -> str:
    return f"{d:%A} {d.day} {d:%B %Y}" if d else ""


def filename(debates: list, late: bool = False) -> str:
    """"Local Government, Housing and Planning Committee, 1 October 2026.docx"
    for one meeting; "Senedd debate summaries 29 September 2026.docx" for
    a morning's email covering several."""
    meetings = {(d.record.forum, d.meeting_date) for d in debates}
    dates = sorted(x for _f, x in meetings if x)
    when = f"{dates[-1].day} {dates[-1]:%B %Y}" if dates else ""
    if len(meetings) == 1:
        forum = next(iter(meetings))[0]
        name = f"{forum}, {when}" if when else forum
    else:
        name = f"Senedd debate summaries {when}".strip()
    return re.sub(r'[\\/:*?"<>|]+', " ", name).strip() + ".docx"


def _witnesses(debate) -> list[tuple[str, str]]:
    """The non-Members who spoke in this item, with their posts, in the
    order they first spoke."""
    out, seen = [], set()
    posts = getattr(debate.record, "attendees", {}) or {}
    for c in debate.item.contributions:
        if c.member or c.speaker in seen:
            continue
        post = posts.get(c.speaker) or c.role
        if not post and not posts:
            continue
        seen.add(c.speaker)
        out.append((c.speaker, post))
    return out


def _speaker(c) -> str:
    return f"{c.speaker} MS" if c.member else c.speaker


def build_debates_document(debates: list, late: bool = False,
                           page_url: str = "") -> bytes | None:
    """The .docx as bytes, or None when there is nothing to detail."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    debates = [d for d in debates if any(d.points)]
    if not debates:
        return None
    ai = any(d.mode == "ai" for d in debates)
    d = _Doc()
    title_text = filename(debates, late)[:-5]

    sec = d.doc.sections[0]
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    d.muted(hp, "NRLA  ·  Senedd debate summaries  ·  " + title_text, size=8.5)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    d.muted(fp, "Page ", size=8.5)
    _page_number(fp)

    t = d.doc.add_paragraph(style="Title")
    t.add_run(title_text)
    sub = d.para(after=2)
    r = sub.add_run("Record now published  ·  in detail" if late else "What was said  ·  in detail")
    r.bold, r.font.size, r.font.color.rgb = True, _pt(12), _rgb(ORANGE)
    org = d.para(after=10)
    d.muted(org, "National Residential Landlords Association", size=10)
    _border_bottom(org)

    intro = d.para(after=8)
    intro.add_run(
        ("This goes with the email of the same name. For each item on NRLA issues "
         "it gives the witnesses and every contribution in turn, summarised by AI "
         "(Claude) from the Senedd's draft Record of Proceedings, which is not yet "
         "final. Every figure is checked against what the speaker said; where a "
         "summary could not be used, the speaker's own sentences are shown in "
         "quotation marks. Use the Record, linked under each heading, when "
         "quoting anyone." if ai else
         "This goes with the email of the same name. For each item on NRLA issues "
         "it gives the witnesses and the most relevant sentences of every "
         "contribution, in the speakers' own words, from the Senedd's draft Record "
         "of Proceedings, which is not yet final.")).font.size = _pt(10)

    meetings: list[tuple] = []
    for deb in debates:
        key = (deb.record.meeting_id, deb.record.forum, deb.meeting_date)
        if key not in meetings:
            meetings.append(key)

    for mid, forum, when in meetings:
        items = [x for x in debates if x.record.meeting_id == mid]
        d.heading(f"{forum}, {_day(when)}" if when else forum, 1, url=items[0].record.url)
        p = d.para(after=6)
        links = [("Record", items[0].record.url)]
        if items[0].papers_url:
            links.append(("Agenda and papers", items[0].papers_url))
        for n, (label, url) in enumerate(links):
            if n:
                d.muted(p, "  ·  ")
            _hyperlink(p, url, label, size=9)

        for deb in items:
            d.heading(deb.title, 2, url=deb.url)
            meta = d.para(after=4)
            d.muted(meta, "The whole item" if deb.whole else
                    f"{len(deb.exchanges)} relevant exchange"
                    f"{'s' if len(deb.exchanges) != 1 else ''} from this item")
            if deb.video_url:
                d.muted(meta, "  ·  ")
                _hyperlink(meta, deb.video_url, "Watch on Senedd.tv", size=9)

            wit = _witnesses(deb)
            if wit:
                w = d.para(after=2)
                rr = w.add_run("Witnesses")
                rr.bold, rr.font.size, rr.font.color.rgb = True, _pt(10), _rgb(DARK_BLUE)
                for name, post in wit:
                    b = d.bullet()
                    rb = b.add_run(name)
                    rb.bold = True
                    rb.font.size = _pt(10)
                    if post:
                        d.muted(b, f", {post}", size=10)

            if deb.overview:
                ov = d.para(after=6)
                ov.add_run(deb.overview).italic = True

            last = None
            for ex, pts in zip(deb.exchanges, deb.points):
                if not pts:
                    continue
                if ex.heading and not deb.whole and ex.heading != last:
                    d.heading(ex.heading, 3)
                    last = ex.heading
                for pt in pts:
                    c = pt.contribution
                    who = d.para(after=1)
                    who.paragraph_format.space_before = _pt(6)
                    who.paragraph_format.keep_with_next = True
                    rw = who.add_run(_speaker(c))
                    rw.bold = True
                    rw.font.color.rgb = _rgb(DARK_BLUE)
                    if c.role:
                        d.muted(who, f"  {c.role}")
                    text = tidy(pt.detail or pt.summary)
                    body = d.para(after=3, indent_cm=0.35)
                    if pt.verbatim:
                        body.add_run(f"“{text}”").italic = True
                    else:
                        body.add_run(text)
                    _border_left(body)
                    anchor = f"{deb.record.url}#{c.anchor}" if c.anchor else ""
                    if anchor:
                        d.muted(body, "  ")
                        _hyperlink(body, anchor, "Record ›", size=8.5, colour=MUTED,
                                   underline=False)

    end = d.para(after=0)
    end.paragraph_format.space_before = _pt(18)
    _border_bottom(end, colour="D9E2E8", size=6)
    note = d.para(size=8.5, colour=MUTED)
    note.add_run(
        ("Summaries are written by AI (Claude) from the Senedd's draft Record of "
         "Proceedings and are not quotations. " if ai else
         "Every word attributed to a speaker is from the Senedd's draft Record of "
         "Proceedings. ")
        + "Contains Senedd Cymru information licensed under the Open Government "
        "Licence v3.0.").font.size = _pt(8.5)
    for run in note.runs:
        run.font.color.rgb = _rgb(MUTED)
    if page_url:
        p = d.para(after=0)
        d.muted(p, "Everything else said in the Chamber and in committee is on the ",
                size=8.5)
        _hyperlink(p, page_url, "live page", size=8.5)
        d.muted(p, ".", size=8.5)

    _fix_settings(d.doc)
    buf = io.BytesIO()
    d.doc.save(buf)
    return buf.getvalue()


__all__ = ["build_debates_document", "filename", "OFF_BLACK"]
